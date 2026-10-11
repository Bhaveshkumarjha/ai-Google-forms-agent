"""Extract plain text from uploaded PDF / DOCX / TXT documents."""
import logging

from filestore import DOCUMENT_EXTS, UploadError, read_validated

log = logging.getLogger(__name__)

MAX_PDF_PAGES = 150
MAX_CHARS = 40000  # keeps the prompt inside free-tier token limits


def extract_text(file_storage, max_bytes):
    """Return (text, truncated). Raises UploadError with a user-safe message."""
    ext, buf = read_validated(file_storage, DOCUMENT_EXTS, max_bytes)
    try:
        if ext == ".txt":
            text = buf.read().decode("utf-8-sig", errors="replace")
        elif ext == ".docx":
            text = _docx_text(buf)
        else:
            text = _pdf_text(buf)
    except UploadError:
        raise
    except Exception:  # corrupt or unusual files must not crash the request
        log.exception("document parsing failed")
        raise UploadError("Could not read this document. Try another file or paste the text.")

    text = text.replace("\x00", "").strip()
    if not text:
        raise UploadError(
            "No text found. Scanned or image-only documents are not supported yet; "
            "paste the questions instead."
        )
    truncated = len(text) > MAX_CHARS
    return text[:MAX_CHARS], truncated


def _docx_text(buf):
    from docx import Document

    doc = Document(buf)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _pdf_text(buf):
    from pypdf import PdfReader

    reader = PdfReader(buf)
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise UploadError("This PDF is password-protected.")
        except UploadError:
            raise
        except Exception:
            raise UploadError("This PDF is password-protected.")
    pages = reader.pages[:MAX_PDF_PAGES]
    return "\n".join((page.extract_text() or "") for page in pages)
